"""Read-only generation checks plus isolated invalid-input fixtures."""
import importlib.util
import tempfile
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("generator", root / "generate-dps-timing-data.py")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
source = root / "trickcal_skillmotion.xlsx"
data = module.build_data(source, dict(module.DEFAULT_SHEETS))
for action in ("lowSkill", "highSkill"):
    obj = data["apostles"]["momo"]["actions"][action]["generatedObjects"][0]
    assert obj["attackSpeedBase"] == 150
    assert obj["attackSpeedReference"] == "生成物自身"
    assert obj["repeatIntervalFrames"] == 121
    assert obj["executionMode"] == "召喚ユニット"
    motion = obj["summonActions"][0]
    assert motion["motionName"] == "Attack1_1"
    assert motion["motionFrames"] == 58
    assert motion["motionSource"] == {"value": 58, "unit": "ゲームF", "gameFrames": 58}
    assert motion["researchStatus"] == "暫定"
    assert "attackMotionFrames" not in obj
    assert not obj["executionIssues"]
    starts = [row for row in obj["timingEvents"]
              if row["recordPurpose"] == "実行" and row["eventType"] == "行動開始"]
    assert len(starts) == (4 if action == "lowSkill" else 1)
    assert all(row["frame"] == 0 and row["researchStatus"] == "暫定" for row in starts)
    observations = [row for row in obj["timingEvents"] if row["recordPurpose"] == "観測"]
    assert len(observations) == (4 if action == "lowSkill" else 1)
    assert all(row["frame"] is not None and row["repeatTarget"] is False for row in observations)
    hits = [row for row in obj["timingEvents"]
            if row["actionId"] and row["recordPurpose"] == "実行" and row["eventType"] == "攻撃"]
    assert len(hits) == 1
    assert hits[0]["frame"] == 20
    assert hits[0]["timeOrigin"] == "召喚ユニット行動開始"
    assert hits[0]["researchStatus"] == "暫定"
assert data["summary"]["summonActions"] == 2
assert sum(row["frame"] is None for action in ("lowSkill", "highSkill")
           for row in data["apostles"]["momo"]["actions"][action]["generatedObjects"][0]["timingEvents"]
           if row["actionId"] and row["recordPurpose"] == "実行") == 0

with tempfile.TemporaryDirectory() as directory:
    book = load_workbook(source)
    sheet = book["生成物基礎設定"]
    headers = [c.value for c in sheet[1]]
    index = headers.index("攻撃速度基礎") + 1
    path = Path(directory) / "fixture.xlsx"
    for value in (None, 0, -1, "invalid", "nan", "inf", True):
        sheet.cell(4, index).value = value
        book.save(path)
        try:
            module.build_data(path, dict(module.DEFAULT_SHEETS))
            raise AssertionError(f"invalid base accepted: {value!r}")
        except ValueError as error:
            assert "Momo_low_clone" in str(error) and "攻撃速度基礎" in str(error)
    sheet.cell(4, index).value = 150
    motion_index = headers.index("攻撃モーション値") + 1
    unit_index = headers.index("攻撃モーション単位") + 1
    for value, unit in ((-1, "ゲームF"), ("invalid", "ゲームF"), ("nan", "ゲームF"),
                        ("inf", "ゲームF"), (True, "ゲームF"), (58, None), (58, "unknown")):
        sheet.cell(4, motion_index).value = value
        sheet.cell(4, unit_index).value = unit
        book.save(path)
        try:
            module.build_data(path, dict(module.DEFAULT_SHEETS))
            raise AssertionError(f"invalid motion accepted: {value!r}/{unit!r}")
        except ValueError as error:
            assert "Momo_low_clone" in str(error) and "攻撃モーション" in str(error)
    for value, unit, expected in ((None, "ゲームF", None), (0, "ゲームF", 0), (1, "ゲーム秒", 60)):
        sheet.cell(4, motion_index).value = value
        sheet.cell(4, unit_index).value = unit
        book.save(path)
        obj = module.build_data(path, dict(module.DEFAULT_SHEETS))["apostles"]["momo"]["actions"]["lowSkill"]["generatedObjects"][0]
        assert obj["attackMotionFrames"] == expected
        assert obj["repeatIntervalFrames"] == 121
    book.close()
print("Generated attack-speed base / inserted column / invalid-input checks passed")
